from pathlib import Path
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"F:\WLWKFSX\QimoProject")
ASSETS = ROOT / "report_assets"
OUT_DOCX = ROOT / "期末综合项目_智能门禁访客管理系统_完成版_扩充版.docx"
FONT = r"C:\Windows\Fonts\simsun.ttc"
FONT_HEI = r"C:\Windows\Fonts\simhei.ttf"


def set_east_asia_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def configure_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_para(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_east_asia_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_east_asia_font(r)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[col_idx])
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_east_asia_font(r, size=9, bold=(row_idx == 0))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
    style_table(table, widths)
    doc.add_paragraph()
    return table


def add_picture(doc, path, caption, width=6.1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_east_asia_font(r, size=9, color="555555")


def make_diagrams():
    ASSETS.mkdir(exist_ok=True)
    font_title = ImageFont.truetype(FONT_HEI, 42)
    font_h = ImageFont.truetype(FONT_HEI, 28)
    font_body = ImageFont.truetype(FONT, 24)
    font_small = ImageFont.truetype(FONT, 20)

    arch = ASSETS / "expanded_architecture.png"
    img = Image.new("RGB", (1600, 900), "#F6F8FB")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "智能门禁访客管理系统总体架构", font=font_title, fill="#16324F")
    layers = [
        (120, 160, 1480, 330, "#EAF3FF", "感知/执行层：CC2530 终端节点"),
        (120, 380, 1480, 550, "#EEF8F0", "网络传输层：ZigBee + 协调器 + 智云平台"),
        (120, 600, 1480, 770, "#FFF6E8", "应用管理层：Python/PyQt5 + 数据库"),
    ]
    for x1, y1, x2, y2, fill, label in layers:
        d.rounded_rectangle((x1, y1, x2, y2), 18, fill=fill, outline="#9AAFC4", width=3)
        d.text((x1 + 24, y1 + 18), label, font=font_h, fill="#1F3A5F")
    boxes = [
        (210, 235, 440, 310, "sensor-a\n温湿度/光照"),
        (500, 235, 730, 310, "sensor-c\nPIR/门磁/门铃"),
        (790, 235, 1020, 310, "sensor-b\n门锁/蜂鸣器/RGB"),
        (1080, 235, 1370, 310, "Coordinator\nZigBee 协调器"),
        (210, 455, 510, 530, "ZXBee 帧协议\nAA...CHK...55"),
        (590, 455, 890, 530, "ZigBee 无线组网\n终端与协调器通信"),
        (970, 455, 1300, 530, "智云平台\nWebSocket/TCP 数据转发"),
        (220, 675, 520, 750, "PyQt5 上位机\n实时看板/远程控制"),
        (620, 675, 920, 750, "业务逻辑\n告警联动/停留计时"),
        (1020, 675, 1320, 750, "数据库\n事件日志/环境数据"),
    ]
    for x1, y1, x2, y2, text in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), 14, fill="#FFFFFF", outline="#6B8AA5", width=2)
        lines = text.split("\n")
        h = len(lines) * 28
        for idx, line in enumerate(lines):
            bbox = d.textbbox((0, 0), line, font=font_body)
            d.text((x1 + (x2 - x1 - bbox[2]) / 2, y1 + (y2 - y1 - h) / 2 + idx * 32), line, font=font_body, fill="#243447")
    img.save(arch)

    flow = ASSETS / "expanded_flow.png"
    img = Image.new("RGB", (1600, 760), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "核心业务流程：感知、判断、联动、记录", font=font_title, fill="#16324F")
    steps = [
        ("访客靠近", "PIR=1\n开始停留计时"),
        ("按门铃/门磁变化", "tch/door 状态上报\n生成事件日志"),
        ("业务判断", "白天/夜间模式\n停留时长/告警等级"),
        ("执行联动", "继电器、蜂鸣器\nRGB 声光提示"),
        ("上位机记录", "看板刷新\n数据库保存与统计"),
    ]
    x = 80
    for i, (title, body) in enumerate(steps):
        d.rounded_rectangle((x, 210, x + 245, 420), 22, fill="#F5F9FF", outline="#79A7D3", width=3)
        d.text((x + 45, 240), title, font=font_h, fill="#1F4D78")
        d.multiline_text((x + 35, 305), body, font=font_body, fill="#334155", spacing=8)
        if i < len(steps) - 1:
            d.line((x + 245, 315, x + 300, 315), fill="#2563EB", width=5)
            d.polygon([(x + 300, 315), (x + 280, 303), (x + 280, 327)], fill="#2563EB")
        x += 300
    d.rounded_rectangle((120, 535, 1480, 650), 18, fill="#FFF7E8", outline="#E6B85C", width=2)
    d.text((150, 565), "说明：我完成了从下位机采集、ZigBee/智云传输、Python 解析、PyQt 展示、远程控制到数据库统计的完整闭环。",
           font=font_small, fill="#7A4F00")
    img.save(flow)
    return arch, flow


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("期末综合项目课程设计报告")
    set_east_asia_font(r, "黑体", 22, True, "1F4D78")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智能门禁访客管理系统")
    set_east_asia_font(r, "黑体", 26, True, "0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("基于 CC2530、ZigBee、智云平台与 PyQt5 的综合物联网系统实现")
    set_east_asia_font(r, "宋体", 12, False, "555555")

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("项目名称", "智能门禁访客管理系统"),
        ("完成情况", "本人完成整个项目的硬件节点、通信协议、上位机、数据库与联调测试"),
        ("课程名称", "智能物联网开发实训"),
        ("学生姓名", "________________"),
        ("学号", "________________"),
        ("完成日期", datetime.now().strftime("%Y年%m月%d日")),
    ]
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
    style_table(table, [2200, 7000])
    doc.add_page_break()


def build():
    arch, flow = make_diagrams()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "摘要", 1)
    add_para(doc, "本报告围绕“智能门禁访客管理系统”的设计与实现展开。项目以 CC2530 和 Z-Stack 为下位机基础，以 ZigBee 作为局域无线通信方式，以智云平台作为数据转发通道，以 Python/PyQt5 作为上位机应用层，并结合数据库完成访客事件和环境数据的持久化管理。")
    add_para(doc, "本项目由本人完成整体方案设计、硬件节点功能实现、通信字段整理、上位机界面开发、数据库设计、联动逻辑调试和报告整理。系统已经形成完整闭环：sensor-a 采集温湿度和光照，sensor-c 负责门禁安防检测，sensor-b 执行门锁、蜂鸣器和 RGB 灯控制，协调器与智云平台完成数据转发，上位机完成实时展示、远程控制、事件记录、统计分析和导出。")
    add_para(doc, "关键词：智能门禁；CC2530；ZigBee；智云平台；PyQt5；数据库；访客管理")

    add_heading(doc, "1 项目完成情况概述", 1)
    add_heading(doc, "1.1 我完成的整体工作", 2)
    add_para(doc, "本次期末综合项目不是单一模块演示，而是一个完整的物联网系统。我完成了从需求分析、模块划分、下位机程序、通信协议、上位机页面、数据库存储到联调测试的全过程。系统可以在上位机中看到门禁状态、访客靠近、门铃、门锁、安全模式、火焰、气体、红外光栅、停留时长、夜间模式、温湿度、光照和运行诊断等信息，并能通过按钮远程控制 sensor-b 与 sensor-c 对应模块。")
    add_para(doc, "最终系统具备“看得见、控得住、查得到、能报警”的完整功能。看得见是指上位机实时显示各节点上报状态；控得住是指远程开门、手动关门、蜂鸣器、告警、解除告警、夜间模式均可下发；查得到是指事件日志和环境数据写入数据库并进行统计；能报警是指夜间靠近、长时间停留和异常状态可以触发声光告警。")
    add_bullets(doc, [
        "硬件部分：完成 sensor-a、sensor-b、sensor-c 三类节点的功能划分与接线调试。",
        "通信部分：统一 ZigBee 节点地址、智云平台 MAC 地址、ZXBee 帧格式和 JSON 字段。",
        "上位机部分：完成浅色 PyQt5 页面、实时数据展示、远程控制、运行诊断、事件时间线和统计图表。",
        "数据库部分：完成访客事件表和环境数据表设计，实现新增、查询、统计和导出。",
        "联调部分：完成开门/关门、门铃短响、告警长响、夜间模式、停留计时、RGB 指示等演示逻辑。",
    ])

    add_heading(doc, "1.2 主要功能清单", 2)
    add_table(doc, ["功能模块", "完成内容", "实现位置/节点"], [
        ["环境采集", "温度、湿度、光照采集并上报，上位机曲线显示", "sensor-a + 上位机环境信息面板"],
        ["门禁检测", "PIR 人体靠近、门磁状态、Touch 门铃、停留时长、夜间模式", "sensor-c + 门禁状态看板"],
        ["远程执行", "远程开门、手动关门、蜂鸣器短响/长响、RGB 颜色指示", "sensor-b + 远程控制区"],
        ["告警联动", "夜间有人靠近、异常开门、长时间停留触发报警", "sensor-c 判断 + sensor-b 执行"],
        ["数据管理", "access_log 和 door_env 两张表记录事件与环境数据", "数据库模块"],
        ["可视化", "状态卡片、实时曲线、事件时间线、运行诊断、访客统计", "PyQt5 上位机"],
    ], [1700, 5000, 2700])

    add_heading(doc, "2 系统需求与总体设计", 1)
    add_heading(doc, "2.1 需求分析", 2)
    add_para(doc, "项目 10 要求完成一个智能门禁访客管理系统，需要能够对门口环境与访客行为进行感知，能够进行门禁控制和异常报警，并能在上位机端进行数据展示和管理。结合课程总体规范，本项目需要体现物联网三层架构、传感器采集、ZigBee 通信、云平台接入、上位机应用和数据库操作。")
    add_para(doc, "我将需求拆分为六类：第一类是采集需求，包含温湿度、光照、人体靠近、门铃、门磁和安防扩展传感器；第二类是执行需求，包含门锁继电器、蜂鸣器和 RGB 灯；第三类是通信需求，要求各节点通过统一协议上报与接收命令；第四类是上位机需求，要求页面清晰、功能可演示；第五类是数据库需求，要求事件和环境数据可保存；第六类是答辩演示需求，要求功能项不能虚假展示，页面只保留已经实现的控制命令。")
    add_heading(doc, "2.2 总体架构", 2)
    add_para(doc, "系统采用“感知执行层、网络传输层、应用管理层”的三层结构。感知执行层由 sensor-a、sensor-b、sensor-c 组成；网络传输层由 Z-Stack、ZXBee 协议、协调器和智云平台组成；应用管理层由 Python/PyQt5 上位机和数据库组成。")
    add_picture(doc, arch, "图 2-1 智能门禁访客管理系统总体架构图", 6.2)
    add_picture(doc, flow, "图 2-2 系统核心业务流程图", 6.2)

    add_heading(doc, "3 硬件节点设计与实现", 1)
    add_heading(doc, "3.1 sensor-a 环境采集节点", 2)
    add_para(doc, "sensor-a 负责环境数据采集，主要完成温度、湿度和光照三个字段的上报。温湿度用于显示门口环境状态，光照用于判断白天/夜间场景，也可以辅助夜间模式展示。该节点定时读取传感器数据并封装为 temp、humi、lux 字段，上位机接收后更新环境信息面板和曲线图。")
    add_heading(doc, "3.2 sensor-b 控制执行节点", 2)
    add_para(doc, "sensor-b 是执行节点，负责真正产生动作效果。它接收上位机下发的 unlock、buzz、rgb、alert、reset 等命令，控制继电器、蜂鸣器和 RGB 灯。为了使演示逻辑清楚，我将开门改为保持开门，不再自动关门；关门必须由上位机手动发送 unlock=0，这样答辩时可以明确说明远程开门和手动关门的区别。")
    add_para(doc, "蜂鸣器逻辑分为短响和长响：Touch 门铃或上位机“触发门铃”使用 buzz=500，表示短响；夜间靠近、停留过久或手动触发告警使用 buzz=1，并配合 alert=2 和红色 RGB，表示持续报警。解除告警使用 reset=1 或 buzz=0，恢复安全状态。")
    add_heading(doc, "3.3 sensor-c 安防检测节点", 2)
    add_para(doc, "sensor-c 负责门禁安防检测，包含 PIR 人体靠近、Touch 门铃、门磁、火焰、可燃气体、红外光栅和夜间模式等状态。PIR 只表示有人靠近，不再被误当成门铃；Touch 才作为门铃短响触发源；夜间模式下检测到靠近或异常状态会触发报警联动。")
    add_para(doc, "上位机端对 sensor-c 上报的 pir 状态增加本地停留计时，避免硬件端 3 秒周期上报导致停留时长不增加的问题。只要 PIR 持续为 1，上位机每秒刷新靠近时长，到达设定阈值后联动 sensor-b 长响报警。")
    add_heading(doc, "3.4 节点功能对应关系", 2)
    add_table(doc, ["节点", "主要字段", "已经实现的功能", "上位机对应位置"], [
        ["sensor-a", "temp、humi、lux", "温湿度和光照采集、曲线显示、环境数据入库", "环境信息面板"],
        ["sensor-b", "unlock、buzz、rgb、alert、reset", "远程开门、手动关门、门铃短响、告警长响、RGB 指示、解除告警", "远程控制区、运行诊断"],
        ["sensor-c", "pir、door、tch、flame、gas、grating、stay、night", "有人靠近、门铃、门磁、火焰/气体/光栅检测、夜间模式、停留计时", "门禁状态看板、事件时间线"],
        ["协调器", "ZXBee 帧", "ZigBee 组网、节点数据转发、云平台桥接", "运行诊断、节点包计数"],
    ], [1200, 2200, 4200, 2100])

    add_heading(doc, "4 通信协议与云平台接入", 1)
    add_para(doc, "通信部分采用 ZigBee 作为短距离无线网络，协调器负责与各终端节点建立通信。业务数据采用 ZXBee 帧封装，负载部分使用 JSON 或 key=value 形式，便于上位机解析。上位机通过 WebSocket 连接智云平台，认证成功后接收 sensor/message 数据包，并根据节点 MAC 判断数据来自 sensor-a、sensor-b 还是 sensor-c。")
    add_table(doc, ["方向", "示例字段", "说明"], [
        ["sensor-a 上行", '{"temp":29.4,"humi":46.4,"lux":457}', "环境数据上报"],
        ["sensor-c 上行", '{"pir":1,"door":1,"tch":1,"alert":1}', "安防状态与门铃事件上报"],
        ["sensor-b 上行", '{"unlock":1,"buzz":1,"rgb":[255,0,0]}', "执行节点状态反馈"],
        ["上位机下行", '{"unlock":0}', "手动关门命令"],
        ["上位机下行", '{"alert":2,"buzz":1,"rgb":[255,0,0]}', "长响告警命令"],
        ["上位机下行", '{"night":1}', "启用夜间模式"],
    ], [1600, 3300, 4500])
    add_para(doc, "调试中我重点处理了字段来源混淆问题：sensor-b 的 unlock 才表示执行节点门锁状态，sensor-c 的 door 表示门磁检测状态。上位机“门状态”卡片最终只跟 unlock 或按钮发出的 unlock 走，避免 sensor-c 周期上报 door=1 把页面错误刷新为已开。")

    add_heading(doc, "5 上位机软件设计与实现", 1)
    add_para(doc, "上位机使用 Python + PyQt5 开发，当前已经改为浅色主题，适合投屏讲解。页面分为标题栏、门禁状态看板、环境信息、远程控制、运行诊断、访客事件时间线和访客统计。每个功能区都对应实际已经实现的命令或数据字段，没有保留无法讲解的虚假按钮。")
    add_picture(doc, ASSETS / "ui_light_dashboard_top.png", "图 5-1 上位机浅色主界面：门禁状态、环境信息与远程控制", 6.25)
    add_picture(doc, ASSETS / "ui_light_dashboard_bottom.png", "图 5-2 上位机浅色主界面：运行诊断、事件时间线与访客统计", 6.25)
    add_heading(doc, "5.1 门禁状态看板", 2)
    add_para(doc, "门禁状态看板负责集中展示系统当前安全状态。访客状态来自 PIR；门状态来自 sensor-b 的 unlock；门铃状态来自 tch；安全模式来自 alert；扩展状态包括火焰、可燃气体、红外光栅、靠近时长和夜间模式。该区域可以在答辩时直接说明每个字段的来源和硬件对应关系。")
    add_heading(doc, "5.2 远程控制区", 2)
    add_para(doc, "远程控制区只展示已经实现的控制命令，包括远程开门、手动关门、触发门铃、关闭蜂鸣器、触发告警、解除告警、启用夜间模式和白天模式。每个按钮都通过统一 command_requested 信号发送字典命令，再由通信层转发给目标节点。")
    add_heading(doc, "5.3 运行诊断与统计", 2)
    add_para(doc, "运行诊断区显示智云连接状态、sensor-a/b/c 包计数、数据库记录数量和最近命令日志。事件时间线以表格方式显示门铃、有人靠近、门打开、告警等事件；统计区可以查看日统计、事件分布、告警趋势，并支持日志与环境数据导出。")

    add_heading(doc, "6 数据库设计", 1)
    add_para(doc, "数据库用于保存系统运行过程中的关键历史数据。项目设计 access_log 和 door_env 两张核心表，分别对应访客事件与环境数据。数据库层封装了插入、查询、删除旧数据、统计日访客量、统计事件类型和导出 CSV 等方法。")
    add_table(doc, ["表名", "字段", "用途"], [
        ["access_log", "id、student_id、sensor、event_type、value、alert、is_remote_unlock、create_time", "保存门铃、PIR、开门、徘徊、报警、远程开门等访客事件"],
        ["door_env", "id、student_id、temp、humi、lux、create_time", "保存温度、湿度、光照等环境历史数据"],
    ], [1500, 5200, 2700])
    add_para(doc, "在上位机界面中，数据库记录数量会显示在运行诊断卡片中；统计图表会从 access_log 中按日期和事件类型查询数据；导出按钮可以将表内容导出为 CSV，便于答辩展示和后续分析。")

    add_heading(doc, "7 关键联动逻辑", 1)
    add_table(doc, ["场景", "判断条件", "系统动作"], [
        ["正常访客靠近", "pir=1，night=0", "页面显示有人靠近，开始停留计时，记录 PIR 事件"],
        ["访客按门铃", "tch=1", "sensor-b 蜂鸣器短响，RGB 蓝色提示，时间线记录门铃"],
        ["远程开门", "上位机发送 unlock=1", "继电器开门并保持，RGB 绿色提示，数据库记录远程开门"],
        ["手动关门", "上位机发送 unlock=0", "继电器关闭，页面门状态改为已关"],
        ["夜间靠近/异常", "night=1 且 pir=1 或 door=1", "发送 alert=2、buzz=1、rgb 红色，触发长响报警"],
        ["解除告警", "reset=1 或 buzz=0", "蜂鸣器关闭，告警状态恢复，页面显示安全"],
    ], [1700, 2900, 4800])
    add_para(doc, "项目调试过程中重点修正了几个容易混淆的问题：第一，PIR 周期上报不是门铃，因此不能让 pir=1 触发短响；第二，Touch 才是门铃短响；第三，夜间靠近应该触发长响报警；第四，开门后不再自动关门，由上位机手动关门；第五，sensor-c 的 door 不能覆盖 sensor-b 的 unlock 门锁状态。")

    add_heading(doc, "8 系统测试与结果", 1)
    add_heading(doc, "8.1 硬件与下位机测试", 2)
    add_table(doc, ["测试项", "测试方法", "结果"], [
        ["PIR 人体检测", "遮挡/靠近 sensor-c 模块，观察 pir 字段和访客状态", "能够显示有人靠近并开始计时"],
        ["Touch 门铃", "按下 Touch 区域，观察 tch 字段和蜂鸣器短响", "按门铃触发短响逻辑"],
        ["继电器门锁", "点击远程开门与手动关门", "unlock=1 开门，unlock=0 关门"],
        ["蜂鸣器", "点击触发门铃、触发告警、关闭蜂鸣器", "短响、长响和关闭逻辑区分清楚"],
        ["RGB 灯", "下发不同 rgb 数组", "可按命令显示对应颜色"],
        ["夜间模式", "启用 night=1 后触发靠近", "上位机联动 sensor-b 长响报警"],
    ], [1700, 4200, 3500])
    add_heading(doc, "8.2 上位机与数据库测试", 2)
    add_para(doc, "上位机测试重点包括字段解析、页面刷新、按钮下发、事件记录和统计展示。通过模拟数据和实际节点上报两种方式进行验证，页面能够实时显示 sensor-a/b/c 的包计数、最新数据和命令日志。事件时间线能够记录门打开、门铃、有人徘徊等事件，数据库统计能够显示事件总数和环境数据总数。")
    add_para(doc, "数据库测试中，access_log 可以保存访客事件，door_env 可以保存环境数据。统计图表使用近 7 天数据生成日统计和事件分布，导出按钮可以将数据导出为 CSV 文件。")

    add_heading(doc, "9 问题处理与优化", 1)
    add_table(doc, ["问题", "原因分析", "处理结果"], [
        ["RGB 灯只有一种颜色", "硬件 RGB 极性与程序假设不一致", "调整 RGB_ON/RGB_OFF 定义，使颜色控制符合实际板卡"],
        ["点击关门后页面又显示开门", "sensor-c 的 door 字段覆盖了 sensor-b 的 unlock 状态", "门状态卡片改为优先使用 unlock，door 只作为安防检测字段"],
        ["PIR 周期上报导致蜂鸣器短响", "把 pir=1 误当成门铃事件", "取消 PIR 触发短响，只有 tch=1 才触发门铃短响"],
        ["停留时长不增加", "硬件周期上报间隔与页面计时逻辑不一致", "上位机增加本地 QTimer，每秒刷新靠近时长"],
        ["夜间靠近无长响", "告警命令只设置 alert，未同时设置 buzz", "夜间报警下发 alert=2、buzz=1、rgb 红色"],
        ["深色页面投屏不清楚", "原主题背景偏暗", "改为浅色主题，图表中文字体显式加载"],
    ], [1800, 3900, 3700])

    add_heading(doc, "10 总结与展望", 1)
    add_para(doc, "本项目完成了一个相对完整的智能门禁访客管理系统。系统覆盖了硬件感知、无线传输、云平台转发、上位机监控、远程控制、数据库管理和统计分析等环节，符合智能物联网开发实训中综合项目对“硬件 + 通信 + 应用 + 数据”的要求。")
    add_para(doc, "通过本项目，我完成了从单点传感器驱动到完整系统联调的全过程，也更加熟悉了 CC2530、Z-Stack、ZigBee 网络、智云平台协议、Python/PyQt5 界面和数据库操作。项目最终能够围绕实际门禁场景进行讲解：访客靠近、按门铃、远程开门、手动关门、夜间报警、解除告警、历史记录和统计分析都可以形成完整演示链路。")
    add_para(doc, "后续如果继续完善，可以增加真实电控锁、摄像头抓拍、人脸识别、手机端推送、权限管理和更长时间的稳定性测试，使系统从课程设计演示进一步接近真实工程应用。")

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
